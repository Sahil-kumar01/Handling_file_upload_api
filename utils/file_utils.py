from fpdf import FPDF
from docx import Document

def is_pdf(file_name: str) -> bool:
    return file_name.lower().endswith('.pdf')

def convert_to_pdf(input_path: str, output_path: str):
    pdf = FPDF()
    pdf.add_page()
    
    if input_path.lower().endswith(('.png', '.jpg', '.jpeg')):
        pdf.image(input_path, x=10, y=10, w=100)
    elif input_path.lower().endswith('.docx'):
        doc = Document(input_path)
        pdf.set_font("Arial", size=12)
        for paragraph in doc.paragraphs:
            safe_text = paragraph.text.encode('latin-1', errors='replace').decode('latin-1')
            pdf.multi_cell(0, 10, txt=safe_text)
    else:
        with open(input_path, 'r', encoding='utf-8', errors='replace') as file:
            text = file.read()
        pdf.set_font("Arial", size=12)
        for line in text.split('\n'):
            safe_line = line.encode('latin-1', errors='replace').decode('latin-1')
            pdf.cell(200, 10, txt=safe_line, ln=True)
    
    pdf.output(output_path)
